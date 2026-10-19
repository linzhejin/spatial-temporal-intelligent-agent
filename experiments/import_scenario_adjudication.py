"""Import a completed human-review workbook into a new development dataset.

The importer never edits its source dataset. It binds the review workbook to
the exact draft bytes and requires a frozen input fingerprint for the output.
"""
from __future__ import annotations

import argparse
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook

from experiments.scenario_evaluator import validate_dataset


_SHA256 = re.compile(r"[0-9a-f]{64}\Z")
_SOURCE_HASH = re.compile(
    r"scenarios_dev_v1\.json\s*\|\s*SHA-256\s+([0-9a-f]{64})"
)
_CASE_HEADERS = {
    "场景ID", "源快照标记", "审查状态", "结论摘要", "审查人", "审查方法", "证据/参考"
}
_TURN_HEADERS = {
    "场景ID", "轮次", "源快照标记", "本轮用户原话", "审查状态",
    "预期响应状态", "预期输出类型", "任务判读", "可行性", "起点 POI ID", "终点 POI ID",
    "途经 POI ID（有序）", "出行方式", "路线策略", "硬约束 JSON",
    "必需输出（JSON列表）", "候选 POI ID（JSON有序列表）", "澄清槽位",
    "允许未知项（JSON列表）", "意图标签", "可接受工具调用序列（JSON二维数组）",
    "事件快照/证据ID", "证据来源/参考", "审查人",
    "理由与备注",
}


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _header_map(sheet) -> dict[str, int]:
    for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        headers = {str(value).strip(): index for index, value in enumerate(values)
                   if value is not None and str(value).strip()}
        if "场景ID" in headers:
            return headers
    raise ValueError(f"{sheet.title}: header row not found")


def _records(sheet, required_headers: set[str]) -> list[dict[str, Any]]:
    headers = _header_map(sheet)
    missing = required_headers - headers.keys()
    if missing:
        raise ValueError(f"{sheet.title}: missing columns: {', '.join(sorted(missing))}")
    records = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        record = {name: values[index] if index < len(values) else None
                  for name, index in headers.items()}
        if not any(value is not None for value in record.values()):
            continue
        records.append(record)
    return records


def _required_text(value: Any, field: str, case_id: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{case_id}: {field} is required")
    return value.strip()


def _optional_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_json(value: Any, field: str, case_id: str, expected_type: type) -> Any:
    if not isinstance(value, str):
        raise ValueError(f"{case_id}: {field} must contain JSON")
    try:
        parsed = json.loads(value)
    except json.JSONDecodeError as exc:
        raise ValueError(f"{case_id}: invalid JSON in {field}: {exc.msg}") from exc
    if not isinstance(parsed, expected_type):
        kind = "object" if expected_type is dict else "array"
        raise ValueError(f"{case_id}: {field} must be a JSON {kind}")
    return parsed


def _find_source_hash(workbook) -> str:
    if "审查说明" not in workbook.sheetnames:
        raise ValueError("workbook is missing 审查说明 sheet")
    for row in workbook["审查说明"].iter_rows(values_only=True):
        for value in row:
            if isinstance(value, str):
                match = _SOURCE_HASH.search(value)
                if match:
                    return match.group(1)
    raise ValueError("workbook is missing the source dataset SHA-256")


def _iso_date(value: Any, case_id: str) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str) and value.strip():
        return value.strip()
    raise ValueError(f"{case_id}: review date is required")


def import_review_workbook(
    *,
    source_path: str | Path,
    workbook_path: str | Path,
    output_path: str | Path,
    snapshot_id: str,
) -> dict[str, Any]:
    """Validate a completed review workbook and save a new adjudicated dev set."""
    source_path = Path(source_path).resolve()
    workbook_path = Path(workbook_path).resolve()
    output_path = Path(output_path).resolve()
    if not _SHA256.fullmatch(snapshot_id):
        raise ValueError("snapshot_id must be a 64-character lowercase SHA-256 fingerprint")
    if output_path in {source_path, workbook_path}:
        raise ValueError("output_path must be different from the source and workbook")
    if output_path.exists():
        raise FileExistsError(f"refusing to overwrite existing output: {output_path}")

    source_bytes = source_path.read_bytes()
    source_hash = _sha256(source_bytes)
    dataset = json.loads(source_bytes.decode("utf-8"))
    validate_dataset(dataset)
    workbook_bytes = workbook_path.read_bytes()
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        if _find_source_hash(workbook) != source_hash:
            raise ValueError("source dataset changed after this workbook was generated")
        if "场景审查" not in workbook.sheetnames or "逐轮裁决" not in workbook.sheetnames:
            raise ValueError("workbook must contain 场景审查 and 逐轮裁决 sheets")
        case_rows = _records(workbook["场景审查"], _CASE_HEADERS)
        turn_rows = _records(workbook["逐轮裁决"], _TURN_HEADERS)
    finally:
        workbook.close()

    source_cases = {case["id"]: case for case in dataset["cases"]}
    reviewed_cases: dict[str, dict[str, Any]] = {}
    for row in case_rows:
        case_id = row.get("场景ID")
        if not isinstance(case_id, str) or case_id not in source_cases:
            raise ValueError(f"unknown or invalid case ID in 场景审查: {case_id!r}")
        if case_id in reviewed_cases:
            raise ValueError(f"duplicate case review row: {case_id}")
        if row.get("审查状态") != "已审":
            raise ValueError(f"{case_id}: case is unreviewed or deferred")
        reviewer = _required_text(row.get("审查人"), "reviewer", case_id)
        method = _required_text(row.get("审查方法"), "review method", case_id)
        if method != "manual_independent":
            raise ValueError(f"{case_id}: review method must be manual_independent")
        evidence = _required_text(row.get("证据/参考"), "case evidence", case_id)
        summary = _required_text(row.get("结论摘要"), "case review summary", case_id)
        rationale = _optional_text(row.get("理由与备注"))
        if row.get("源快照标记") != source_cases[case_id].get("snapshot_id"):
            raise ValueError(f"{case_id}: source snapshot marker was changed")
        review_evidence = [f"case conclusion: {summary}", evidence]
        if rationale:
            review_evidence.append(f"case rationale: {rationale}")
        reviewed_cases[case_id] = {
            "reviewer": reviewer,
            "evidence": review_evidence,
            "turns": {},
        }

    if reviewed_cases.keys() != source_cases.keys():
        missing = sorted(source_cases.keys() - reviewed_cases.keys())
        extra = sorted(reviewed_cases.keys() - source_cases.keys())
        raise ValueError(f"case review coverage mismatch; missing={missing}, extra={extra}")

    for row in turn_rows:
        case_id = row.get("场景ID")
        if not isinstance(case_id, str) or case_id not in source_cases:
            raise ValueError(f"unknown or invalid case ID in 逐轮裁决: {case_id!r}")
        case = source_cases[case_id]
        turns = reviewed_cases[case_id]["turns"]
        turn_number = row.get("轮次")
        if type(turn_number) is not int or turn_number < 1:
            raise ValueError(f"{case_id}: turn number must be a positive integer")
        if turn_number in turns:
            raise ValueError(f"{case_id}: duplicate turn {turn_number}")
        if turn_number > len(case["messages"]):
            raise ValueError(f"{case_id}: turn number exceeds source conversation")
        if row.get("审查状态") != "已审":
            raise ValueError(f"{case_id} turn {turn_number}: unreviewed turn")
        if row.get("源快照标记") != case.get("snapshot_id"):
            raise ValueError(f"{case_id} turn {turn_number}: source snapshot marker was changed")
        if row.get("本轮用户原话") != case["messages"][turn_number - 1]["content"]:
            raise ValueError(f"{case_id} turn {turn_number}: user message was changed")

        reviewer = _required_text(row.get("审查人"), "reviewer", case_id)
        if reviewer != reviewed_cases[case_id]["reviewer"]:
            raise ValueError(f"{case_id} turn {turn_number}: reviewer differs from case reviewer")
        task_interpretation = _required_text(row.get("任务判读"), "task interpretation", case_id)
        rationale = _required_text(row.get("理由与备注"), "review rationale", case_id)
        evidence = _optional_text(row.get("证据来源/参考"))
        event_evidence = _optional_text(row.get("事件快照/证据ID"))
        if not evidence and not event_evidence:
            raise ValueError(f"{case_id} turn {turn_number}: evidence is required")
        if evidence:
            reviewed_cases[case_id]["evidence"].append(f"turn {turn_number}: {evidence}")
        reviewed_cases[case_id]["evidence"].append(
            f"turn {turn_number} task: {task_interpretation}"
        )
        reviewed_cases[case_id]["evidence"].append(
            f"turn {turn_number} rationale: {rationale}"
        )
        if event_evidence:
            reviewed_cases[case_id]["evidence"].append(
                f"turn {turn_number} event evidence: {event_evidence}"
            )

        gold = {
            "status": _required_text(row.get("预期响应状态"), "expected status", case_id),
            "kind": _required_text(row.get("预期输出类型"), "expected kind", case_id),
            "feasibility": _required_text(row.get("可行性"), "feasibility", case_id),
            "start_poi_id": _optional_text(row.get("起点 POI ID")),
            "end_poi_id": _optional_text(row.get("终点 POI ID")),
            "via_poi_ids": _parse_json(row.get("途经 POI ID（有序）"), "途经 POI ID", case_id, list),
            "travel_mode": _optional_text(row.get("出行方式")),
            "strategy": _optional_text(row.get("路线策略")),
            "hard_constraints": _parse_json(row.get("硬约束 JSON"), "hard constraints", case_id, dict),
            "required_outputs": _parse_json(row.get("必需输出（JSON列表）"), "required outputs", case_id, list),
            "candidate_poi_ids": _parse_json(row.get("候选 POI ID（JSON有序列表）"), "candidate POI IDs", case_id, list),
            "clarification_slot": _optional_text(row.get("澄清槽位")),
            "allowed_unknowns": _parse_json(row.get("允许未知项（JSON列表）"), "allowed unknowns", case_id, list),
            "intent_label": _required_text(row.get("意图标签"), "intent_label", case_id),
            "expected_tools": _parse_json(
                row.get("可接受工具调用序列（JSON二维数组）"), "expected_tools", case_id, list
            ),
        }
        turns[turn_number] = (gold, _iso_date(row.get("审查日期"), case_id), rationale)

    for case_id, case in source_cases.items():
        reviewed = reviewed_cases[case_id]
        if set(reviewed["turns"]) != set(range(1, len(case["messages"]) + 1)):
            raise ValueError(f"{case_id}: reviewed turn coverage does not match source messages")
        case["expected_turns"] = [reviewed["turns"][i][0]
                                  for i in range(1, len(case["messages"]) + 1)]
        case["snapshot_id"] = snapshot_id
        case["adjudication"] = {
            "status": "adjudicated",
            "method": "manual_independent",
            "reviewer": reviewed["reviewer"],
            "evidence": list(dict.fromkeys(reviewed["evidence"])),
        }

    dataset["authoring_status"] = (
        "Independently adjudicated development dataset; not a formal test split."
    )
    dataset["adjudication_import"] = {
        "source_dataset_sha256": source_hash,
        "review_workbook_sha256": _sha256(workbook_bytes),
        "snapshot_id": snapshot_id,
    }
    validate_dataset(dataset)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    serialized = json.dumps(dataset, ensure_ascii=False, indent=2) + "\n"
    with output_path.open("x", encoding="utf-8", newline="\n") as stream:
        stream.write(serialized)
    return dataset


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True,
                        help="the exact scenarios_dev_v1.json used to create the workbook")
    parser.add_argument("--workbook", type=Path, required=True,
                        help="completed scenario adjudication workbook")
    parser.add_argument("--output", type=Path, required=True,
                        help="new output JSON path; existing files are never overwritten")
    parser.add_argument("--snapshot-id", required=True,
                        help="64-character fingerprint from the frozen research input inventory")
    args = parser.parse_args(argv)
    imported = import_review_workbook(
        source_path=args.source,
        workbook_path=args.workbook,
        output_path=args.output,
        snapshot_id=args.snapshot_id,
    )
    print(json.dumps({
        "output": str(args.output.resolve()),
        "cases": len(imported["cases"]),
        "turns": sum(len(case["expected_turns"]) for case in imported["cases"]),
        "snapshot_id": args.snapshot_id,
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
