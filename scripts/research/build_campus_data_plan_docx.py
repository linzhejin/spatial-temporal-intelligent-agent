"""Build an editable campus data verification and acquisition work plan.

Inputs are read-only research snapshots. No project routing or POI data is changed.
"""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


NAVY = "243B53"
PALE = "F3F7FA"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
FONT = "Microsoft YaHei"


def apply_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = FONT
    run.font.color.rgb = BLACK
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), FONT)


def set_cell_fill(cell, color: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top: int = 90, start: int = 105, bottom: int = 90, end: int = 105) -> None:
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = margins.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}")
            margins.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tblpr = table._tbl.tblPr
    borders = tblpr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), BORDER)


def set_repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    trpr.append(marker)


def set_no_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def add_hyperlink(paragraph, label: str, url: str) -> None:
    rel_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), FONT)
    props.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "26547C")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def paragraph(doc, text: str = "", *, bold_lead: str | None = None, style: str | None = None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        apply_font(p.add_run(bold_lead), bold=True)
        apply_font(p.add_run(text[len(bold_lead):]))
    else:
        apply_font(p.add_run(text))
    return p


def heading(doc, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def table(doc, headers: list[str], rows: list[list[str]], widths: list[float], *, font_size: float = 9.3):
    result = doc.add_table(rows=1, cols=len(headers))
    result.autofit = False
    set_table_borders(result)
    for col, width in zip(result.columns, widths, strict=True):
        col.width = Inches(width)
    for idx, value in enumerate(headers):
        cell = result.rows[0].cells[idx]
        cell.width = Inches(widths[idx])
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell, 120, 105, 120, 105)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(value)
        apply_font(run, font_size, True)
        run.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_header(result.rows[0])
    set_no_split(result.rows[0])
    for row_idx, row_data in enumerate(rows):
        row = result.add_row()
        set_no_split(row)
        for cell_idx, value in enumerate(row_data):
            cell = row.cells[cell_idx]
            cell.width = Inches(widths[cell_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_idx % 2:
                set_cell_fill(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.13
            apply_font(p.add_run(value), font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return result


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.top_margin = Inches(0.76)
    section.bottom_margin = Inches(0.70)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.7)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.26
    for name, size, before, after in (("Title", 19, 0, 12), ("Heading 1", 13.5, 16, 6), ("Heading 2", 11.4, 10, 4)):
        style = styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.font.bold = name != "Title"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        if name == "Title":
            style.paragraph_format.line_spacing = 1.1
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for key in ("ascii", "hAnsi", "eastAsia"):
            fonts.set(qn(f"w:{key}"), FONT)
        if name == "Title":
            ppr = style.element.get_or_add_pPr()
            border = ppr.find(qn("w:pBdr"))
            if border is not None:
                ppr.remove(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    apply_font(footer.add_run("校园数据核实规划  |  2026年10月9日  |  "), 8.0)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for item in (begin, code, end):
        run = OxmlElement("w:r")
        run.append(item)
        footer._p.append(run)


def load_inputs(source_root: Path, queue_path: Path, poi_path: Path):
    manifest = json.loads((source_root / "research/campus_enrichment/source_manifest.json").read_text(encoding="utf-8"))
    with queue_path.open("r", encoding="utf-8-sig", newline="") as handle:
        raw_queue = list(csv.DictReader(handle))
    counts = Counter(row["filter_reason"] for row in raw_queue)
    expected = {"access_node": 34, "access_way": 2, "campus_polygon": 33}
    if len(raw_queue) != 69 or counts != expected:
        raise ValueError(f"The 69-item OSM review snapshot changed: {len(raw_queue)} rows, {counts}")
    if any(row["verification_status"] != "pending_field_or_official_review" for row in raw_queue):
        raise ValueError("Some OSM review items have changed status; revise the document narrative first")
    if any(not row["source_link"].startswith("https://www.openstreetmap.org/") for row in raw_queue):
        raise ValueError("An OSM object link is missing or unexpected")
    poi = json.loads(poi_path.read_text(encoding="utf-8"))
    if poi.get("count") != len(poi.get("pois", [])):
        raise ValueError("POI count and POI array disagree")
    return manifest, raw_queue, len(poi["pois"])


def add_task_table(doc: Document, data: list[tuple[str, str, str, str, str]], *, future: bool = False) -> None:
    headers = ["编号", "事项", "已知证据与缺口", "下一步行动与获取对象", "完成标准与跟踪"]
    rows = [[code, title, evidence, action, done + "\n状态：待核实；结果：待填；复核：待定"] for code, title, evidence, action, done in data]
    table(doc, headers, rows, [0.45, 0.93, 1.87, 2.12, 1.55], font_size=8.65 if future else 8.85)


def add_sources(doc: Document, source_root: Path, queue_path: Path) -> None:
    heading(doc, "资料依据")
    paragraph(doc, "数据截至2026年10月9日研究快照，更新时应重新核对来源、数量和现行校方规则。")
    paragraph(doc, "独立研究工作树：06_武大校园专属步行数据增强规划_20261009.md、07_武大步行通行事实现场核验清单_20261009.md、research/campus_enrichment/source_manifest.json。")
    paragraph(doc, f"69项原始核验队列：{queue_path.name}（本机D盘05_reports目录）。")
    p = doc.add_paragraph()
    apply_font(p.add_run("官方网页："), 9.4)
    for index, (label, url) in enumerate([
        ("2025年新版预约系统公告", "https://news.whu.edu.cn/info/1013/488247.htm"),
        ("2026年校园交通治理协作报道", "https://news.whu.edu.cn/info/2301/501027.htm"),
        ("2026年暑期校巴服务报道", "https://news.whu.edu.cn/info/1013/501127.htm"),
    ]):
        add_hyperlink(p, label, url)
        if index < 2:
            apply_font(p.add_run("  ·  "), 9.4)


def build_document(manifest: dict, queue: list[dict], poi_count: int, source_root: Path, queue_path: Path) -> Document:
    doc = Document()
    configure(doc)
    doc.add_paragraph("武汉大学校园步行导航数据核实与后续获取规划", style="Title")
    paragraph(doc, "版本 1.0  |  2026年10月9日  |  项目持续修订用")
    paragraph(doc, "本规划用于维护校园步行导航的数据质量：先把来源、许可和真实通行条件说清，再核实关键校门、入口与路段，最后通过独立路线对照判断新增数据是否改善导航。任务表和附录保留状态与复核栏，核验后可逐项更新。")
    paragraph(doc, "当前最需要解决的是校门与道路的现行通行规则、建筑真实入口，以及混合路网和历史校园 GIS 的来源与使用条件。遥感和公开地图已能提出候选路线，尚不足以证明这些路在现场可走、舒适或适合无障碍通行。")

    heading(doc, "一 现有数据与实际可用程度")
    paragraph(doc, "判定口径：取得表示文件或记录已到手；处理表示完成了提取、匹配或计算；核实表示有可追溯的规则或现场证据；正式使用还需要与具体导航场景相符的来源、许可、有效期和独立验证。以下按2026年10月9日研究快照列示。")
    inventory = [
        ["正式路网", "研究快照的混合路网为EPSG:4326、14,622条有向边；边属性有版本/哈希记录。", "可作现有最短路和候选属性的对照骨架。", "校园边界、来源许可、局部连接和步行权限未逐段核实。"],
        ["现有POI", f"本次只读检查项目data/pois.json为{poi_count}条；名称、别名和坐标可读取。", "可供目的地检索及路线起终点吸附。", "数量会随另一数据工作变化；实际入口、别名准确性和路网连接仍需抽样核实。"],
        ["OSM通行与入口", "Geofabrik湖北PBF快照为2026-10-06；提取2,994条候选道路及559个出入口、障碍和过街节点，形成69项断点队列。", "可定位值得复核的门点、限制标签、台阶及连通缺口。", "原始标签和OSM学部范围不是现行校方政策；队列69项全部待核验。"],
        ["CHMv2树冠", "Meta/WRI CHMv2瓦片已取得，并按路段周边计算树冠高度代理；绑定边ID与几何哈希。", "筛选可能有遮荫差异的路线。", "瓦片未找到观测日期；树冠高度不等于当下行人阴凉或舒适度。"],
        ["SinoLC地表分类", "SinoLC-1武汉成员已取得，ZIP成员CRC32、长度及SHA-256已核对；母包MD5未验证。", "辅助比较植被与地表类型候选。", "影像时间混合且较旧；NoData与分类不代表真实通行条件。"],
        ["SinoBF建筑功能", "SinoBF-1 V1.0武汉图层已提取，建筑功能仅用于候选筛查。", "查找可能遗漏的楼宇功能。", "缺官方名称、入口和开放时间；CC BY-NC 4.0用途需按场景审查。"],
        ["Overture建筑与地点", "2026-09-23.1版已按研究范围作本地筛选；Buildings及Places保留不同来源许可。", "辅助楼宇轮廓与地点查漏。", "存在分类错配或重复候选；轮廓中心不是入口，混合来源需保留逐条授权信息。"],
    ]
    table(doc, ["数据", "已取得与处理", "可支持的功能", "尚缺的证明"], inventory, [0.95, 2.45, 1.55, 1.97], font_size=8.75)
    paragraph(doc, "以上数量对应数据快照，不代表现行校方全量数据。正式路线是否安全、开放及按身份可通行，仍以有效规则与获准现场核验为准。")

    heading(doc, "二 当前需要核实的事项")
    paragraph(doc, "按对路线正确性的影响安排核实。V01至V03先处理范围、许可和通行约束；V04至V06核实目的地、无障碍与环境；V07处理历史GIS资料。每项默认状态为待核实，结果和下次复核日期留给后续填写。")
    verification = [
        ("V01", "校园边界", "现有图网范围和三个OSM学部关系可用于筛选；33项断点由项目多边形裁剪产生。项目范围尚无校方确认。", "向项目数据负责人或校方确认研究边界及边界版本；逐项判定33处道路是否属于目标区域。", "边界依据、版本及33项范围判定可追溯；不把裁剪写成物理封闭。"),
        ("V02", "路网来源", "混合路网含课程图资；ArcEngine-WGIS等候选的许可、采集日期与坐标系有疑点。", "核对原始资料、许可范围、CRS、采集年代和拓扑语义；请导师或持有者确认书面条件。", "研究、公开展示、衍生分发条件分别记录；争议图层留在候选层。"),
        ("V03", "校门与道路通行", "69项队列含34个access_node及2个access_way；2025预约系统上线已确认，具体门/时段仍未知。", "先向保卫部门或导师引荐团队核实身份、进出方向、开放时间、预约前提和临时封闭；获准后核对门点。", "每条规则有对象、适用人群/方式、方向、有效期、来源和复核人；冲突标为待裁定。"),
        ("V04", "楼宇与入口", f"当前POI快照有{poi_count}条，SinoBF与Overture仅提供候选名称/轮廓。", "选常用教学楼、图书馆、医院及校门，核对别名、实际步行入口和入口到道路的连接。", "抽样点能到达正确入口；建筑中心不得充当已核实的门。"),
        ("V05", "无障碍与路面", "OSM有台阶、坡度与路面标签候选，完整性和现场现况未知。", "在获准区域记录台阶、坡道、路缘、路面与前后连续性，优先覆盖关键OD差异路段。", "逐段标明观测日期、位置、通行方式和无法确认的环节；形成可复查记录。"),
        ("V06", "环境偏好", "CHMv2和SinoLC仅为代理；od_04、od_02的离线计算仍未现场验证。", "同一天时段比较同图网最短路与候选路，记录天气、树荫、实际距离、用时和是否抵达入口。", "至少一组独立实走对照有原始记录；未观察到收益就记无收益或不确定。"),
        ("V07", "历史校园GIS", "ArcEngine-WGIS有人行道、站点和线路候选；BikeMap有静态道路点边；两者数据权利和时相未明。", "向作者或导师核实数据文件的许可、底图来源、采集日期、CRS和字段含义；披露已做的本机临时只读审查。", "取得明确许可与元数据后再评估是否值得独立比对；未经确认不并入正式路网。"),
    ]
    add_task_table(doc, verification)
    heading(doc, "重点门点的首轮核查")
    gate_note = paragraph(doc, "相近点仅用于安排同一次到访，不能据此认定为同一个校门。下表的最大受影响分量是图论影响指标，不代表标签错误或现场封闭。")
    gate_note.paragraph_format.keep_with_next = True
    clusters = [
        ["凌波门", "1199819110、13613667177、13613667184", "foot=no与foot=permit并存；一项标注07:00至22:00。核对各开口及方向。"],
        ["文澜门", "13613667193、1663906491、13613667196", "foot=no与foot=permit并存；核对全天标注是否仍适用。"],
        ["弘毅门", "13241555638、13613667157", "foot=no与foot=permit并存；核对预约、身份及实际位置。"],
        ["珞珈门近邻簇", "13238420133、13238420137、13240918484、13613667148", "foot=no与foot=permit并存；逐点核对进出方向与通行条件。"],
    ]
    table(doc, ["门点候选", "原始OSM节点ID", "现场或校方需回答的问题"], clusters, [1.1, 2.75, 3.07], font_size=9.1)

    heading(doc, "三 后续值得获取的数据")
    paragraph(doc, "获取顺序按步行导航价值、现有缺口和获取可行性安排。优先向校方索取窄范围、带有效期的静态或脱敏事实；拿不到导出时，可先取得具体规则确认与获准核验渠道。")
    future = [
        ("D01", "当前通行规则", "当前门点队列与历史公告给出线索；规则没有身份组和周期时段的可靠现况。", "保卫部及校园交通协作团队：校门/通道ID、坐标、人员类型、方向、开放时段、预约要求、临时封闭起止与维护人。", "能按时间和身份判定一条路线是否可走；每条附来源、有效期与更新方式。"),
        ("D02", "真实入口与常用设施", "现有POI可检索；楼宇轮廓和质心无法代表入口。", "校方地图/院系/现场：楼名与别名、设施类别、入口坐标、入口对应路段、开放时段、更新时间。", "抽样目的地导航到正确入口，名称及坐标可追溯。"),
        ("D03", "无障碍连续性", "OSM标签稀疏；台阶、坡道与路缘的连续路线未获证实。", "校方无障碍设施资料或获准现场观察：台阶级数、坡道、路缘、宽度/障碍、通行方向及变动日期。", "关键路线从起点到目标入口连续可达；未知段显式标记。"),
        ("D04", "缺失人行连接", "现有图可能遗漏穿楼通道、支路和校园内部连接；历史GIS仅是候选。", "获准的路网图层或定点踏勘：端点、线形、连接方式、是否常开、限制、采集日期、CRS和许可。", "候选边与正式图拓扑匹配并经实地核对，断点变化可解释。"),
        ("D05", "阴凉与地面体验", "遥感树冠与地表是代理，无法说明不同时间的地面体感。", "分时段现场观察：路段ID、天气、时间、阴凉范围、路面、步行时间、观测方式。", "固定样本可复测；与最短路比较时说明真实绕行和体验差异。"),
        ("D06", "静态校巴资料", "2026年暑期约15分钟班次是当时报道，不能代表学期常规服务。", "后勤运输责任部门：线路几何、站点坐标与顺序、运营日、班次/时段、有效日期、更新渠道和许可。", "可独立核对一条常规学期线路；接驳评估保留纯步行方案。"),
        ("D07", "路线偏好样本", "现有离线OD对照少且没有地面真值，不能据此训练学习式路线决策。", "在已获准、可通行的路线中收集匿名成对选择与任务完成结果：OD、路线版本、时间天气、到达入口、用时、反馈。", "样本可复核且覆盖不同场景；只用作后续偏好研究，不保存个人轨迹或身份。"),
    ]
    add_task_table(doc, future, future=True)
    heading(doc, "低优先级和独立扩展")
    paragraph(doc, "WHU-Mix Vector仅在确认存在珞珈山校区覆盖、逐源许可和清楚影像日期后，再比较其建筑轮廓是否补足现有候选；它本身不提供入口。i2Nav-Robot适合定位或建图评测，只有研究目标扩展到这一方向才评估下载。校巴实时到站数据另作后续扩展，需正式授权的只读接口和故障回退，不以小程序私有接口作为来源。")

    heading(doc, "四 执行安排与路线对照")
    table(doc, ["阶段", "先做什么", "完成后进入下一阶段的依据"], [
        ["第一阶段", "确认研究边界、路网/历史GIS来源与许可；向校方核实行人规则的获得方式。", "范围版本明确；可进入现场的区域与需联系部门明确；未知规则保持未知。"],
        ["第二阶段", "分流69项：先审33项边界，再核34个门点和2条道路限制；核常用入口与台阶坡道。", "关键对象有独立证据、日期、适用对象和冲突记录；不凭图论结果放宽限制。"],
        ["第三阶段", "在经确认可走的图网上，独立走路线并比较最短路与环境偏好路。", "至少一组OD记录正确入口、通行可行性、实际距离/时间、阴凉及失败原因。"],
    ], [1.05, 2.95, 2.92], font_size=9.05)
    paragraph(doc, "首轮两组路线仅用于决定哪些差异值得核验。两张图网必须分别计算、分别记录；路线边缺少显式foot/access不表示可以通行。")
    table(doc, ["路线候选", "离线发现", "外业前提及记录"], [
        ["樱花大道至珞珈山 od_04", "β=0.25树冠代理：混合图+11.6个百分点、绕行+1.84%；OSM-base图+10.5个百分点、绕行+1.50%。", "先确认使用图网及其许可，再检查山路、台阶、开放条件与终点入口；此路线不作无障碍样本。"],
        ["总馆至工学分馆 od_02", "混合图树冠代理+6.3个百分点、绕行+0.88%；OSM-base图树冠偏好未换路。", "核对两端实际入口与差异路段；该对照主要用于判断结论对图网来源有多敏感。"],
    ], [1.48, 2.55, 2.89], font_size=9.05)
    paragraph(doc, "当前路况事件机制能表示绑定到边的一次性临时管制及明确的生效、失效时刻；身份组、每周重复开放的校门规则需要在取得授权数据后单独设计。结果未核实、发生冲突或超过有效期时，记录为未知并安排复核。")

    heading(doc, "五 记录规则与进入正式数据的条件")
    paragraph(doc, "每个任务保留固定编号、当前证据、缺口、下一步行动、获取对象、完成标准、状态、核验结果和下次复核日期。状态推荐使用待核实、已确认开放、已确认关闭、有条件通行、证据冲突、不适用。任何状态变更要附证据链接或获准的现场记录编号、核验日期与复核人角色。")
    paragraph(doc, "数据取得后先逐项记录是否允许内部研究、公开展示、派生使用和再分发，以及坐标系、更新责任方和有效期。只有许可和通行事实均明确、关键路线独立验证通过，才考虑形成正式数据变更；环境偏好要有可理解的地面收益，并保留最短路对照。")
    add_sources(doc, source_root, queue_path)
    doc.add_page_break()
    heading(doc, "附录一 69项连通性核验索引")
    paragraph(doc, "索引按入口节点、道路限制、边界裁剪分组。编号A01至A69为本文档的稳定工作编号；原CSV行号方便回查原文件。每行保留OSM对象链接、坐标和原核验问题。当前69项均为待校方或现场核实，后续状态须由记录模板补证据后更新。")
    ordered = sorted(enumerate(queue, start=2), key=lambda pair: ({"access_node": 0, "access_way": 1, "campus_polygon": 2}[pair[1]["filter_reason"]], pair[0]))
    labels = {"access_node": "入口节点", "access_way": "道路限制", "campus_polygon": "边界裁剪"}
    headers = ["编号", "原行", "类型", "OSM对象", "WGS84坐标", "原核验问题", "状态"]
    appendix = table(doc, headers, [], [0.42, 0.40, 0.79, 1.23, 1.22, 2.12, 0.74], font_size=8.2)
    for idx, (csv_line, item) in enumerate(ordered, start=1):
        row = appendix.add_row()
        set_no_split(row)
        vals = [f"A{idx:02d}", str(csv_line), labels[item["filter_reason"]], "", f'{item["lat"]},\n{item["lon"]}', item["verification_question"], "待核实"]
        for col, value in enumerate(vals):
            cell = row.cells[col]
            cell.width = Inches([0.42, 0.40, 0.79, 1.23, 1.22, 2.12, 0.74][col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 75, 85, 75, 85)
            if idx % 2 == 0:
                set_cell_fill(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.10
            if col == 3:
                add_hyperlink(p, f'{item["object_type"]}/{item["osm_id"]}', item["source_link"])
            else:
                apply_font(p.add_run(value), 7.95)

    heading(doc, "附录二 单项核验记录模板")
    paragraph(doc, "每次核验复制下面的字段填写一条记录；若同一行涉及多个开口或不同时间条件，分别记录，保留原始对象与历史状态。")
    table(doc, ["字段", "填写内容"], [
        ["任务编号与对象", "V或D任务编号；如关联69项，另填A编号、原CSV行、OSM对象链接。"],
        ["当前证据与缺口", "现有来源、版本、原标签、尚未确认的问题。"],
        ["核验结果", "开放／关闭／有条件／冲突／不适用／未核实；写明观察到的实体与路网连接。"],
        ["适用条件", "适用人员与方式、进出方向、工作日和假日时段、预约前提。"],
        ["空间与时间", "坐标及CRS、定位误差、核验时间、有效起止、下次复核日期。"],
        ["证据与责任", "官方通知URL或文号、获准现场记录编号、核验人角色、更新责任部门。"],
        ["许可与下一步", "内部研究、公开展示、派生与分发的许可范围；冲突时需向谁确认。"],
    ], [1.3, 5.62], font_size=9.2)
    heading(doc, "附录三 修订记录")
    table(doc, ["版本", "日期", "变更内容", "核验人"], [["1.0", "2026年10月9日", "依据现有研究快照建立任务、69项索引和记录口径。", "待填写"], ["下一版", "待填写", "待填写", "待填写"]], [0.75, 1.35, 3.85, 0.97], font_size=9.2)
    doc.core_properties.title = "武汉大学校园步行导航数据核实与后续获取规划"
    doc.core_properties.subject = "校园导航数据核验与持续修订"
    return doc


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--queue", type=Path, required=True)
    parser.add_argument("--poi-file", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    manifest, queue, poi_count = load_inputs(args.source_root, args.queue, args.poi_file)
    doc = build_document(manifest, queue, poi_count, args.source_root, args.queue)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(json.dumps({"output": str(args.output), "queue_rows": len(queue), "poi_count": poi_count, "task_rows": 14}, ensure_ascii=False))


if __name__ == "__main__":
    main()
